"""Blueprint TG: consolida a planilha do Tesouro Gerencial (dados/tg/) no modelo ORC usado pelas abas Orçamento e Execução.

  GET  /api/tg/orcamento
  POST /api/tg/atualizar_join   roda scripts/criar_tg_execucao_join.py (botão API-TG do RADAR.html)
  GET  /api/tg/planejamento     tabela tg_planejamento, previsão por tipo e mês (botão PLANEJAMENTO, ao lado de INSTRUMENTOS)

Fonte dos arquivos de dados/tg/: pasta do Google Drive "03.001.1.Radar". Desde 03/10/2026 o download é
AUTOMÁTICO pela API do Google Drive (radar_backend/tg_drive.py): ao abrir o RADAR e no botão API-TG, só os
arquivos novos/alterados; credencial OAuth (App para computador) e token em ~/.radar_fnac.
  https://drive.google.com/drive/folders/1SX6fLRByOnCEE-i_LL-Jg8a4t-n4nnPR
"""

from __future__ import annotations

import csv
import os
import re
import sqlite3
import sys
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

from flask import Blueprint, jsonify, request

from radar_backend.radar_config import INSTRUMENTOS_DB, SCRIPTS_DIR, TG_DIR
from radar_backend.radar_util import (
    _tg_data_iso,
    agora_iso,
    escrever_log,
    resposta_erro,
)

bp = Blueprint("tg_orcamento", __name__)

DRIVE_TG_URL = "https://drive.google.com/drive/folders/1SX6fLRByOnCEE-i_LL-Jg8a4t-n4nnPR"   # pasta "03.001.1.Radar"


@bp.post("/api/tg/atualizar_join")
def api_tg_atualizar_join():
    """Roda scripts/criar_tg_execucao_join.py (relê dados/tg/ + as tabelas transferegov_ted_*
    já coletadas e reconstrói tg_execucao_join). Rápido (não baixa nada), roda em linha."""
    try:
        if str(SCRIPTS_DIR) not in sys.path:
            sys.path.insert(0, str(SCRIPTS_DIR))
        import importlib
        import criar_tg_execucao_join as _mod
        importlib.reload(_mod)   # garante que uma edição no script seja pega sem reiniciar o backend
        resumo = _mod.executar()
        escrever_log(f"TG | tg_execucao_join atualizada: {resumo['linhas_apos_dedupe']} linhas, "
                     f"cobertura {resumo['cobertura']}/{resumo['total_instrumentos']}, "
                     f"{len(resumo['avisos'])} aviso(s)")
        return jsonify({"sucesso": True, **resumo})
    except Exception as exc:
        return resposta_erro(exc)


@bp.get("/api/tg/planejamento")
def api_tg_planejamento():
    """PLANEJAMENTO (tabela tg_planejamento, db_bi_previsao.csv — ver scripts/criar_tg_execucao_join.py): previsão por
    instrumento SEPARADA POR TIPO — Previsão total de desembolsos, Pagamentos do ano, Sobras/Faltas de crédito, RAP a pagar e,
    por mês, Desembolso / Pagamentos / Financeiro. Tipos diferentes nunca são somados entre si (usuário, 28/09/2026)."""
    try:
        con = sqlite3.connect(INSTRUMENTOS_DB)
        con.row_factory = sqlite3.Row
        if not con.execute("SELECT 1 FROM sqlite_master WHERE name='tg_planejamento'").fetchone():
            return jsonify({"sucesso": True, "instrumentos": [], "aviso": "tg_planejamento ainda não existe — rode API-TG."})
        linhas = con.execute("""
            SELECT p.processo_sei, i.tipo_instrumento, i.localidades, i.valor_atual,
                   p.termo, p.pi, p.gnd, p.item, p.mes, p.valor
            FROM tg_planejamento p
            JOIN instrumentos i ON i.processo_sei = p.processo_sei
            ORDER BY i.tipo_instrumento, p.processo_sei, p.mes
        """).fetchall()
        con.close()
        MENSAIS = {"Desembolso", "Pagamentos", "Financeiro"}
        agrupado: dict[str, dict] = {}
        anos = set()
        for r in linhas:
            g = agrupado.setdefault(r["processo_sei"], {
                "processo_sei": r["processo_sei"], "tipo_instrumento": r["tipo_instrumento"],
                "localidades": r["localidades"], "valor_atual": r["valor_atual"],
                "previsao_total": None, "pagamentos_ano": None, "sobras_faltas": None, "rap_a_pagar": None,
                "mensal": {k: {} for k in MENSAIS}, "linhas": {},
            })
            item, v = r["item"], r["valor"]
            chave = {"RAP A Pagar": "rap_a_pagar", "Sobras (+) ou Faltas (-) de Crédito": "sobras_faltas"}.get(item)
            if item.startswith("Previsão Total de Desembolsos"):
                chave = "previsao_total"
                anos.add(item[-4:])
            elif item.startswith("Pagamentos ") and item[-4:].isdigit():
                chave = "pagamentos_ano"
            if chave:
                g[chave] = (g[chave] or 0) + v
            elif item in MENSAIS and r["mes"]:
                g["mensal"][item][r["mes"]] = g["mensal"][item].get(r["mes"], 0) + v
            # origem: cada linha da planilha (Termo / PI / GND), para conferência
            o = g["linhas"].setdefault((r["termo"], r["pi"], r["gnd"]), {"termo": r["termo"], "pi": r["pi"], "gnd": r["gnd"], "itens": []})
            o["itens"].append({"item": item, "mes": r["mes"], "valor": v})
        instrumentos = []
        for g in agrupado.values():
            g["linhas"] = list(g["linhas"].values())
            instrumentos.append(g)
        instrumentos.sort(key=lambda g: -(g["previsao_total"] or 0))
        return jsonify({"sucesso": True, "instrumentos": instrumentos, "total": len(instrumentos),
                        "ano": max(anos) if anos else None, "fonte": "db_bi_previsao.csv (aba Previsão da Programação)"})
    except Exception as exc:
        return resposta_erro(exc)


# ============================================================
# TG -> BASE ORÇAMENTÁRIA DO RADAR  (abas Orçamento / Execução)
# ============================================================

_TG_EXPORT_HINT = re.compile(r"tes[_ ]?gerencial|tesouro|bi[_ ]?tes|db[_ ]?bi", re.I)


def _tg_arquivo_mais_recente() -> Path | None:
    """Planilha TG mais recente. Prioriza nomes de export do Tesouro Gerencial
    (db_bi_tesgerencial…) para não confundir com outras planilhas soltas na pasta."""
    TG_DIR.mkdir(parents=True, exist_ok=True)
    from radar_backend import tg_arquivos                 # pasta apagada -> recriada da cópia guardada no Instrumentos.db
    tg_arquivos.restaurar()
    candidatos = []
    for padrao in ("*.xlsx", "*.xlsm", "*.csv"):
        candidatos.extend(TG_DIR.glob(padrao))
    candidatos = [p for p in candidatos if p.is_file() and not p.name.startswith("~$")]
    if not candidatos:
        return None
    exports = [p for p in candidatos if _TG_EXPORT_HINT.search(p.name)]
    alvo = exports or candidatos
    return max(alvo, key=lambda p: p.stat().st_mtime)


TG_ITENS_ORC = {
    9: 'Dotação Inicial',
    13: 'Dotação Atualizada',
    18: 'Destaque Concedido',
    20: 'Crédito Indisponível',
    23: 'Despesas Empenhadas',
    28: 'Despesas Pagas',
    50: 'RAP Inscrito',
    51: 'RAP Cancelado',
    52: 'RAP Pago',
    53: 'RAP a Pagar',
}


def _tg_money(v) -> float:
    if v is None or v == '':
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace('R$', '').replace(' ', '')
    if not s or s.upper() == 'NA':
        return 0.0
    # Formato brasileiro 1.234,56. Tambem reconhece milhares sem centavos
    # (ex.: 35.000), comuns em exportacoes do Tesouro Gerencial.
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    elif re.fullmatch(r'[+-]?\d{1,3}(?:\.\d{3})+', s):
        s = s.replace('.', '')
    s = re.sub(r'[^0-9+\-.]', '', s)
    try:
        return float(s)
    except Exception:
        return 0.0


def _tg_iter_registros(arquivo: Path, aba_pedida: str = ''):
    """Retorna (abas, aba, colunas, iterator_de_dicts) sem truncar a base."""
    if arquivo.suffix.lower() == '.csv':
        f = arquivo.open('r', encoding='utf-8-sig', errors='replace', newline='')
        amostra = f.read(8192); f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(amostra, delimiters=';,\t,')
        except csv.Error:
            dialect = csv.excel; dialect.delimiter = ';'
        reader = csv.reader(f, dialect)
        try:
            header = next(reader)
        except StopIteration:
            header = []
        cols = [str(x or '').strip() or f'COL_{i+1}' for i, x in enumerate(header)]
        def gen():
            try:
                for row in reader:
                    yield {cols[i]: (row[i] if i < len(row) else '') for i in range(len(cols))}
            finally:
                f.close()
        return ['CSV'], 'CSV', cols, gen()

    from openpyxl import load_workbook
    wb = load_workbook(arquivo, read_only=True, data_only=True)
    abas = wb.sheetnames
    aba = aba_pedida if aba_pedida in abas else (abas[0] if abas else '')
    if not aba:
        wb.close()
        return abas, '', [], iter(())
    ws = wb[aba]
    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        header = ()
    cols=[]; usados={}
    for i,v in enumerate(header):
        base=str(v).strip() if v is not None else ''
        base=base or f'COL_{i+1}'
        usados[base]=usados.get(base,0)+1
        cols.append(base if usados[base]==1 else f'{base}_{usados[base]}')
    def gen():
        try:
            for row in it:
                yield {cols[i]: (row[i] if i < len(row) else None) for i in range(len(cols))}
        finally:
            wb.close()
    return abas, aba, cols, gen()


@bp.get('/api/tg/orcamento')
def api_tg_orcamento():
    """Consolida a planilha TG mais recente no mesmo modelo ORC usado pelo RADAR."""
    try:
        arquivo = _tg_arquivo_mais_recente()
        if not arquivo:
            return jsonify({'sucesso': False, 'erro': 'Nenhuma planilha TG encontrada em dados/tg/.'}), 404
        ano = int(request.args.get('ano', 2026))
        aba_pedida = str(request.args.get('aba') or '').strip()
        abas, aba, colunas, registros = _tg_iter_registros(arquivo, aba_pedida)
        required = ['acao','localizador','fonte','gnd','item_cod','valor','data']
        mapa = {str(c).strip(): c for c in colunas}
        faltantes = [c for c in required if c not in mapa]

        # Se nenhuma aba foi explicitamente escolhida, procura automaticamente
        # a aba que contém a estrutura padrão do relatório do Tesouro Gerencial.
        if faltantes and not aba_pedida and arquivo.suffix.lower() != '.csv':
            try:
                if hasattr(registros, 'close'):
                    registros.close()
            except Exception:
                pass
            for candidata in abas:
                a2, aba2, cols2, reg2 = _tg_iter_registros(arquivo, candidata)
                mapa2 = {str(c).strip(): c for c in cols2}
                falt2 = [c for c in required if c not in mapa2]
                if not falt2:
                    aba, colunas, registros, mapa, faltantes = aba2, cols2, reg2, mapa2, []
                    break
                try:
                    if hasattr(reg2, 'close'):
                        reg2.close()
                except Exception:
                    pass

        if faltantes:
            return jsonify({
                'sucesso': False,
                'erro': 'A planilha TG não possui as colunas exigidas para Orçamento/Execução: ' + ', '.join(faltantes),
                'arquivo': arquivo.name,
                'aba': aba,
                'colunas': colunas,
            }), 422

        grupos = {}; por_pi = {}; processadas=0; total=0; ignoradas_ano=0
        for row in registros:
            total += 1
            data = _tg_data_iso(row.get('data'))
            if not data.startswith(str(ano)):
                if data: ignoradas_ano += 1
                continue
            try:
                item_cod = int(float(str(row.get('item_cod') or '').replace(',', '.')))
            except Exception:
                continue
            if item_cod not in TG_ITENS_ORC:
                continue
            acao=str(row.get('acao') or '').strip(); localizador=str(row.get('localizador') or '').strip()
            fonte=str(row.get('fonte') or '').strip(); gnd=str(row.get('gnd') or '').strip()
            if not acao or not localizador:
                continue
            key='|'.join([acao,localizador,fonte,gnd])
            if key not in grupos:
                grupos[key] = {
                    'acao':acao,'localizador':localizador,'fonte':fonte,'gnd':gnd,
                    'acao_nome':str(row.get('acao_nome') or '').strip(),
                    'localizador_nome':str(row.get('localizador_nome') or '').strip(),
                    'gnd_nome':str(row.get('gnd_nome') or '').strip(),
                    'uge':str(row.get('uge') or '').strip(),'uge_nome':str(row.get('uge_nome') or '').strip(),
                    'uo':str(row.get('uo') or '').strip(),'funcao':str(row.get('funcao') or '').strip(),
                    'subfuncao':str(row.get('subfuncao') or '').strip(),'programa':str(row.get('programa') or '').strip(),
                    'rp':str(row.get('rp') or '').strip(),'rp_nome':str(row.get('rp_nome') or '').strip(),
                    'nd':str(row.get('nd') or '').strip(),'nd_nome':str(row.get('nd_nome') or '').strip(),
                    'ptres':str(row.get('ptres') or '').strip(),'itens':{},'itensPorMes':{},'pos':{},'pis':{}
                }
            g=grupos[key]
            po=str(row.get('po') or '').strip()
            if po: g['pos'][po]=str(row.get('po_nome') or '').strip()
            valor=_tg_money(row.get('valor'))
            ic=str(item_cod)
            g['itens'][ic]=g['itens'].get(ic,0.0)+valor
            mes=data[:7]
            g['itensPorMes'].setdefault(ic,{})
            g['itensPorMes'][ic][mes]=g['itensPorMes'][ic].get(mes,0.0)+valor
            pi=str(row.get('pi') or '').strip()
            if pi:
                g['pis'][pi]=str(row.get('pi_nome') or '').strip()
                if pi not in por_pi:
                    por_pi[pi]={'pi':pi,'acao':acao,'localizador':localizador,'itens':{}}
                por_pi[pi]['itens'][ic]=por_pi[pi]['itens'].get(ic,0.0)+valor
            processadas += 1

        return jsonify({
            'sucesso': True,
            'fonte': 'TG', 'arquivo': arquivo.name, 'caminho': 'dados/tg/'+arquivo.name,
            'modificado_em': datetime.fromtimestamp(arquivo.stat().st_mtime).astimezone().isoformat(timespec='seconds'),
            'abas': abas, 'aba': aba, 'ano': ano, 'rowCount': processadas,
            'totalRowsInFile': total, 'skippedYear': ignoradas_ano,
            'grupos': grupos, 'porPI': por_pi,
        })
    except Exception as exc:
        return resposta_erro(exc)
